#!/usr/bin/env python3
"""
Test Excel file headers to verify they match the review request requirements
"""

import requests
import io
from openpyxl import load_workbook

BACKEND_URL = "https://cash-flow-mgr-1.preview.emergentagent.com/api"

def test_excel_headers():
    session = requests.Session()
    
    # Login
    login_data = {"username": "operaciones@ibsgroup.mx", "password": "Sistema2"}
    response = session.post(f"{BACKEND_URL}/auth/login", json=login_data)
    print(f"Login: {response.status_code}")
    
    print("\n=== Testing Excel File Headers ===")
    
    # Test 1: Cash Excel Export Headers
    print("\n1. Testing Cash Excel Export Headers...")
    response = session.get(f"{BACKEND_URL}/export/efectivo/xlsx")
    if response.status_code == 200:
        # Load Excel file from response
        excel_file = io.BytesIO(response.content)
        wb = load_workbook(excel_file)
        ws = wb.active
        
        # Get headers from first row
        headers = []
        for col in range(1, ws.max_column + 1):
            cell_value = ws.cell(row=1, column=col).value
            if cell_value:
                headers.append(cell_value)
        
        print(f"Found headers: {headers}")
        
        # Expected headers from review request
        expected_headers = [
            "Fecha", "Tipo Movimiento", "Origen/Destino Tipo", "Origen/Destino Nombre",
            "Afectación", "Cantidad", "Folio/Cheque", "Concepto", "Usuario"
        ]
        
        # Check if all expected headers are present
        missing_headers = []
        for expected in expected_headers:
            if expected not in headers:
                missing_headers.append(expected)
        
        if not missing_headers:
            print("✅ All expected headers found in Cash Excel export")
        else:
            print(f"❌ Missing headers in Cash Excel export: {missing_headers}")
    else:
        print(f"❌ Failed to get Cash Excel export: {response.status_code}")
    
    # Test 2: Bank Balances Excel Export Headers
    print("\n2. Testing Bank Balances Excel Export Headers...")
    response = session.get(f"{BACKEND_URL}/export/bank_balances/xlsx")
    if response.status_code == 200:
        # Load Excel file from response
        excel_file = io.BytesIO(response.content)
        wb = load_workbook(excel_file)
        ws = wb.active
        
        # Get headers from first row
        headers = []
        for col in range(1, ws.max_column + 1):
            cell_value = ws.cell(row=1, column=col).value
            if cell_value:
                headers.append(cell_value)
        
        print(f"Found headers: {headers}")
        
        # Expected headers from review request
        expected_headers = ["Nombre Cuenta", "Banco", "Saldo", "Última Actualización"]
        
        # Check if all expected headers are present
        missing_headers = []
        for expected in expected_headers:
            if expected not in headers:
                missing_headers.append(expected)
        
        if not missing_headers:
            print("✅ All expected headers found in Bank Balances Excel export")
        else:
            print(f"❌ Missing headers in Bank Balances Excel export: {missing_headers}")
            
        # Check if there's actual data
        if ws.max_row > 1:
            print(f"✅ Excel file contains {ws.max_row - 1} data rows")
            
            # Check a sample row for timestamp format
            if ws.max_row >= 2:
                timestamp_cell = ws.cell(row=2, column=4).value  # "Última Actualización" column
                print(f"Sample timestamp: {timestamp_cell}")
        else:
            print("⚠️ Excel file contains headers only, no data rows")
    else:
        print(f"❌ Failed to get Bank Balances Excel export: {response.status_code}")

if __name__ == "__main__":
    test_excel_headers()