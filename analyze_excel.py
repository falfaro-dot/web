from openpyxl import load_workbook

# Load the workbook
wb = load_workbook('/app/test_layout.xlsm', data_only=True)
ws = wb.active

print("=== ANÁLISIS DEL ARCHIVO EXCEL ===\n")

# Extract client data from fixed cells
print("1. DATOS DEL CLIENTE (celdas fijas):")
print(f"   D2 (Emisor): {ws['D2'].value}")
print(f"   D6 (Cliente): {ws['D6'].value}")
print(f"   D7 (RFC): {ws['D7'].value}")
print(f"   D8 (Calle): {ws['D8'].value}")
print(f"   D9 (Colonia): {ws['D9'].value}")
print(f"   D10 (Estado): {ws['D10'].value}")

# Extract description from E13
print(f"\n2. DESCRIPCIÓN:")
print(f"   E13: {ws['E13'].value}")

# Search for financial values in column K starting from row 12
print(f"\n3. VALORES FINANCIEROS (Columna K desde fila 12):")
for row in range(12, 50):  # Search up to row 50
    cell_k = ws[f'K{row}'].value
    cell_l = ws[f'L{row}'].value
    
    if cell_k and isinstance(cell_k, str):
        cell_k_lower = cell_k.lower().strip()
        if 'subtotal' in cell_k_lower:
            print(f"   Fila {row} - {cell_k}: {cell_l}")
        elif 'iva' in cell_k_lower:
            print(f"   Fila {row} - {cell_k}: {cell_l}")
        elif 'total' in cell_k_lower and 'factura' in cell_k_lower:
            print(f"   Fila {row} - {cell_k}: {cell_l}")

# Show some additional context
print(f"\n4. ESTRUCTURA GENERAL:")
print(f"   Hoja activa: {ws.title}")
print(f"   Dimensiones: {ws.dimensions}")

# Show header row around row 12
print(f"\n5. ENCABEZADOS (fila 12):")
for col in ['A', 'B', 'C', 'D', 'E', 'F', 'G', 'H', 'I', 'J', 'K', 'L']:
    print(f"   {col}12: {ws[f'{col}12'].value}")

wb.close()
